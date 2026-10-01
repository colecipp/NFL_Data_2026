from pathlib import Path
import re

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ==================================================
# SETTINGS
# ==================================================

SEASON = 2025

OUTPUT_DIR = Path("output")
OUTPUT_DIR.mkdir(exist_ok=True)

MASTER_OUTPUT = f"NFL_Player_Ratings_{SEASON}.xlsx"
MASTER_PARQUET = OUTPUT_DIR / f"NFL_Player_Ratings_{SEASON}.parquet"

# If True, the build stops when any expected position model is missing.
# This protects the master workbook from silently being incomplete.
REQUIRE_ALL_MODELS = True

# If True, the build stops if one Player_ID appears in more than one model.
# This is especially useful for OT/OG-C and EDGE/LB role-classification checks.
FAIL_ON_DUPLICATE_PLAYER_IDS = True

# Keep every position-specific advanced metric in the combined workbook.
# Set False later if you want a compact ratings-only master.
KEEP_ALL_POSITION_METRICS = True


# ==================================================
# POSITION SOURCES
#
# Parquet is preferred because it preserves the exact dataframe schema.
# Excel is only a fallback if the position script did not create Parquet.
# ==================================================

POSITION_SOURCES = [
    {
        "model_group": "QB",
        "parquet": OUTPUT_DIR / f"QB_{SEASON}.parquet",
        "excel": Path(f"QB_Advanced_{SEASON}.xlsx"),
        "sheet": "All_QBs",
    },
    {
        "model_group": "HB",
        "parquet": OUTPUT_DIR / f"HB_{SEASON}.parquet",
        "excel": Path(f"HB_Advanced_{SEASON}.xlsx"),
        "sheet": "All_HBs",
    },
    {
        "model_group": "WR",
        "parquet": OUTPUT_DIR / f"WR_{SEASON}.parquet",
        "excel": Path(f"WR_Advanced_{SEASON}.xlsx"),
        "sheet": "All_WRs",
    },
    {
        "model_group": "TE",
        "parquet": OUTPUT_DIR / f"TE_{SEASON}.parquet",
        "excel": Path(f"TE_Advanced_{SEASON}.xlsx"),
        "sheet": "All_TEs",
    },
    {
        "model_group": "OT",
        "parquet": OUTPUT_DIR / f"OT_{SEASON}.parquet",
        "excel": Path(f"OT_Advanced_{SEASON}.xlsx"),
        "sheet": "All_OTs",
    },
    {
        "model_group": "IOL",
        "parquet": OUTPUT_DIR / f"OG_C_{SEASON}.parquet",
        "excel": Path(f"OG_C_Advanced_{SEASON}.xlsx"),
        "sheet": "All_Interior_OL",
    },
    {
        "model_group": "EDGE",
        "parquet": OUTPUT_DIR / f"DE_EDGE_{SEASON}.parquet",
        "excel": Path(f"DE_EDGE_Advanced_{SEASON}.xlsx"),
        "sheet": "All_EDGE",
    },
    {
        "model_group": "DT",
        "parquet": OUTPUT_DIR / f"DT_{SEASON}.parquet",
        "excel": Path(f"DT_Advanced_{SEASON}.xlsx"),
        "sheet": "All_DTs",
    },
    {
        "model_group": "LB",
        "parquet": OUTPUT_DIR / f"LB_{SEASON}.parquet",
        "excel": Path(f"LB_Advanced_{SEASON}.xlsx"),
        "sheet": "All_LBs",
    },
    {
        "model_group": "CB",
        "parquet": OUTPUT_DIR / f"CB_{SEASON}.parquet",
        "excel": Path(f"CB_Advanced_{SEASON}.xlsx"),
        "sheet": "All_CBs",
    },
    {
        "model_group": "S",
        "parquet": OUTPUT_DIR / f"S_{SEASON}.parquet",
        "excel": Path(f"S_Advanced_{SEASON}.xlsx"),
        "sheet": "All_Safeties",
    },
]


# ==================================================
# MASTER COLUMN ORDER
# ==================================================

CORE_COLUMNS = [
    "Player",
    "Player_ID",
    "Team_2025",
    "Position",
    "Model_Group",
    "Position_Rank",
    "Age",
    "Age_Source",
    "Final_Overall",
    "Age_Adjusted_Overall",
    "Legacy_Rating",
    "Legacy_Rating_Used",
    "Legacy_Fallback_Applied",
    "Advanced_Rating",
    "Advanced_Weighted_Pctl",
    "Pre_Bonus_Overall",
    "Historical_Record_Bonus",
    "Advanced_Data_Coverage",
    "Actual_Advanced_Weight",
    "Actual_Legacy_Weight",
    "Legacy_Match_Status",
    "Madden",
    "PFF",
    "PFR",
    "pfr_id",
    "pff_id",
    "espn_id",
]

COMPACT_COLUMNS = CORE_COLUMNS.copy()


POSITION_SORT_ORDER = {
    "QB": 1,
    "HB": 2,
    "WR": 3,
    "TE": 4,
    "OT": 5,
    "OG": 6,
    "C": 7,
    "EDGE": 8,
    "DE": 8,
    "DT": 9,
    "NT": 9,
    "LB": 10,
    "ILB": 10,
    "MLB": 10,
    "CB": 11,
    "S": 12,
    "FS": 12,
    "SS": 12,
}


# ==================================================
# HELPERS
# ==================================================

def read_excel_fallback(path, preferred_sheet):
    """
    Read a position workbook when its Parquet is unavailable.

    First try the expected sheet name. If that exact sheet does not exist,
    use the first sheet whose name starts with 'All_'.
    """

    excel_file = pd.ExcelFile(path)

    if preferred_sheet in excel_file.sheet_names:
        sheet_name = preferred_sheet
    else:
        all_sheets = [
            sheet
            for sheet in excel_file.sheet_names
            if str(sheet).startswith("All_")
        ]

        if not all_sheets:
            raise ValueError(
                f"{path} exists, but no '{preferred_sheet}' or 'All_*' "
                "sheet was found."
            )

        sheet_name = all_sheets[0]

    return pd.read_excel(
        path,
        sheet_name=sheet_name
    )


def clean_excel_duplicate_headers(df):
    """
    Excel adds suffixes like '.1' when a dataframe originally contained
    repeated column names. Keep the unsuffixed version when both exist.
    """

    columns_to_drop = []

    for column in df.columns:
        column_string = str(column)

        if "." not in column_string:
            continue

        base, suffix = column_string.rsplit(".", 1)

        if (
            suffix.isdigit()
            and base in df.columns
        ):
            columns_to_drop.append(column)

    if columns_to_drop:
        df = df.drop(
            columns=columns_to_drop,
            errors="ignore"
        )

    return df


def load_position_source(source):
    parquet_path = source["parquet"]
    excel_path = source["excel"]
    model_group = source["model_group"]

    if parquet_path.exists():
        df = pd.read_parquet(parquet_path)
        source_used = str(parquet_path)

    elif excel_path.exists():
        df = read_excel_fallback(
            excel_path,
            source["sheet"]
        )

        df = clean_excel_duplicate_headers(df)
        source_used = str(excel_path)

    else:
        return None, None

    df = df.copy()
    df["Model_Group"] = model_group
    df["Model_Source_File"] = source_used

    return df, source_used


def ensure_common_columns(df):
    """
    Make sure every model has the common master fields even if an older
    position script did not yet create one of them.
    """

    common_defaults = {
        "Player": np.nan,
        "Player_ID": np.nan,
        "Team_2025": np.nan,
        "Position": np.nan,
        "Age": np.nan,
        "Age_Source": np.nan,
        "Final_Overall": np.nan,
        "Age_Adjusted_Overall": np.nan,
        "Legacy_Rating": np.nan,
        "Legacy_Rating_Used": np.nan,
        "Legacy_Fallback_Applied": False,
        "Advanced_Rating": np.nan,
        "Advanced_Weighted_Pctl": np.nan,
        "Pre_Bonus_Overall": np.nan,
        "Historical_Record_Bonus": 0.0,
        "Advanced_Data_Coverage": np.nan,
        "Actual_Advanced_Weight": np.nan,
        "Actual_Legacy_Weight": np.nan,
        "Legacy_Match_Status": np.nan,
        "Madden": np.nan,
        "PFF": np.nan,
        "PFR": np.nan,
        "pfr_id": np.nan,
        "pff_id": np.nan,
        "espn_id": np.nan,
    }

    for column, default in common_defaults.items():
        if column not in df.columns:
            df[column] = default

    # Older QB code may have Legacy_Rating but not the newer
    # Legacy_Rating_Used audit field.
    df["Legacy_Rating_Used"] = (
        df["Legacy_Rating_Used"]
        .fillna(df["Legacy_Rating"])
    )

    return df



def normalize_identifier_value(value):
    """
    Normalize IDs imported from mixed Excel/Parquet sources.

    Examples:
      57488      -> "57488"
      57488.0    -> "57488"
      "57488"    -> "57488"
      NaN / None -> <NA>

    PyArrow requires one consistent type per Parquet column, while the
    individual position files can legitimately load the same ID as either
    an integer, float, or string.
    """

    if pd.isna(value):
        return pd.NA

    if isinstance(value, (int, np.integer)):
        return str(int(value))

    if isinstance(value, (float, np.floating)):
        if np.isfinite(value) and float(value).is_integer():
            return str(int(value))

        return str(value)

    value = str(value).strip()

    if value == "" or value.lower() in {
        "nan",
        "none",
        "<na>",
    }:
        return pd.NA

    # Clean strings such as "57488.0" created by an Excel numeric ID column.
    if re.fullmatch(r"-?\d+\.0", value):
        return value[:-2]

    return value


def normalize_identifier_columns(df):
    """
    Force every ID-like column to pandas' nullable string dtype before
    writing the combined master Parquet.

    This prevents ArrowInvalid errors caused by mixing integer IDs from
    Parquet position files with string IDs from Excel fallbacks.
    """

    id_columns = [
        column
        for column in df.columns
        if (
            str(column).lower().endswith("_id")
            or str(column).lower() == "player_id"
        )
    ]

    for column in id_columns:
        df[column] = (
            df[column]
            .map(normalize_identifier_value)
            .astype("string")
        )

    return df


def calculate_position_rank(df):
    """
    Rank players within their model group by current Final_Overall.
    IOL is intentionally one combined OG/C model group.
    """

    df["Position_Rank"] = (
        df.groupby(
            "Model_Group",
            dropna=False
        )["Final_Overall"]
        .rank(
            method="min",
            ascending=False,
            na_option="bottom"
        )
    )

    df["Position_Rank"] = (
        pd.to_numeric(
            df["Position_Rank"],
            errors="coerce"
        )
        .astype("Int64")
    )

    return df


def find_duplicate_players(df):
    valid_ids = (
        df["Player_ID"]
        .notna()
        &
        df["Player_ID"]
        .astype(str)
        .str.strip()
        .ne("")
    )

    duplicate_mask = (
        valid_ids
        &
        df.duplicated(
            "Player_ID",
            keep=False
        )
    )

    duplicate_columns = [
        "Player",
        "Player_ID",
        "Team_2025",
        "Position",
        "Model_Group",
        "Final_Overall",
        "Age_Adjusted_Overall",
        "Model_Source_File",
    ]

    return (
        df.loc[
            duplicate_mask,
            [
                column
                for column in duplicate_columns
                if column in df.columns
            ]
        ]
        .sort_values(
            ["Player_ID", "Model_Group"]
        )
    )


def ordered_master_columns(df):
    """
    Common fields first, then every position-specific metric in the order
    the individual files introduced them.
    """

    front = [
        column
        for column in CORE_COLUMNS
        if column in df.columns
    ]

    remaining = [
        column
        for column in df.columns
        if column not in front
    ]

    return front + remaining


def format_master_workbook(path):
    workbook = load_workbook(path)
    worksheet = workbook["All_Players"]

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    header_lookup = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    # Useful fixed widths for the important identity/rating columns.
    preferred_widths = {
        "Player": 24,
        "Player_ID": 15,
        "Team_2025": 12,
        "Position": 10,
        "Model_Group": 12,
        "Position_Rank": 13,
        "Age": 9,
        "Age_Source": 20,
        "Final_Overall": 14,
        "Age_Adjusted_Overall": 20,
        "Legacy_Rating": 14,
        "Advanced_Rating": 15,
        "Legacy_Match_Status": 24,
        "Model_Source_File": 28,
    }

    for column_name, width in preferred_widths.items():
        if column_name in header_lookup:
            column_letter = get_column_letter(
                header_lookup[column_name]
            )

            worksheet.column_dimensions[
                column_letter
            ].width = width

    # Keep the giant union-of-metrics workbook readable without letting
    # long headers create absurdly wide columns.
    for column_index in range(
        1,
        worksheet.max_column + 1
    ):
        column_letter = get_column_letter(
            column_index
        )

        if (
            worksheet.column_dimensions[
                column_letter
            ].width is not None
        ):
            continue

        header_value = worksheet.cell(
            row=1,
            column=column_index
        ).value

        header_length = len(
            str(header_value)
        ) if header_value is not None else 10

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(header_length + 2, 11),
            24
        )

    # Two-decimal formatting for overall/rating fields.
    for column_name in [
        "Final_Overall",
        "Age_Adjusted_Overall",
        "Legacy_Rating",
        "Legacy_Rating_Used",
        "Advanced_Rating",
        "Advanced_Weighted_Pctl",
        "Pre_Bonus_Overall",
        "Historical_Record_Bonus",
    ]:
        if column_name not in header_lookup:
            continue

        column_index = header_lookup[column_name]

        for row in range(
            2,
            worksheet.max_row + 1
        ):
            worksheet.cell(
                row=row,
                column=column_index
            ).number_format = "0.00"

    # One decimal for age.
    if "Age" in header_lookup:
        age_column = header_lookup["Age"]

        for row in range(
            2,
            worksheet.max_row + 1
        ):
            worksheet.cell(
                row=row,
                column=age_column
            ).number_format = "0.0"

    workbook.save(path)


# ==================================================
# LOAD EVERY POSITION MODEL
# ==================================================

position_frames = []
missing_models = []

print(
    "\nLoading position models..."
)

for source in POSITION_SOURCES:
    frame, source_used = load_position_source(
        source
    )

    if frame is None:
        missing_models.append(
            source["model_group"]
        )

        print(
            f"  MISSING: {source['model_group']}"
        )

        continue

    frame = ensure_common_columns(frame)

    position_frames.append(frame)

    print(
        f"  {source['model_group']}: "
        f"{len(frame)} players "
        f"from {source_used}"
    )


if missing_models:
    message = (
        "\nMissing position models: "
        + ", ".join(missing_models)
    )

    if REQUIRE_ALL_MODELS:
        raise FileNotFoundError(
            message
            +
            "\nRun those position scripts before building the master."
        )

    print(
        "WARNING:"
        + message
    )


if not position_frames:
    raise RuntimeError(
        "No position outputs were found."
    )


# ==================================================
# COMBINE
# ==================================================

master = pd.concat(
    position_frames,
    ignore_index=True,
    sort=False
)


# Remove completely empty rows if an Excel fallback happened to include one.
master = master.dropna(
    how="all",
    subset=[
        "Player",
        "Player_ID",
        "Final_Overall",
    ]
).copy()


# ==================================================
# DUPLICATE PLAYER CHECK
# ==================================================

duplicate_players = find_duplicate_players(
    master
)

if not duplicate_players.empty:
    print(
        "\nDUPLICATE PLAYER IDs FOUND ACROSS POSITION MODELS:"
    )

    print(
        duplicate_players.to_string(
            index=False
        )
    )

    if FAIL_ON_DUPLICATE_PLAYER_IDS:
        raise ValueError(
            "\nThe master build stopped because at least one player "
            "appears in multiple position models. Fix the role classifier "
            "or set FAIL_ON_DUPLICATE_PLAYER_IDS = False if the overlap "
            "is intentional."
        )


# ==================================================
# POSITION RANK / SORT
# ==================================================

master = calculate_position_rank(
    master
)


master["_Position_Sort"] = (
    master["Position"]
    .astype(str)
    .str.upper()
    .str.strip()
    .map(POSITION_SORT_ORDER)
    .fillna(99)
)


master = master.sort_values(
    by=[
        "_Position_Sort",
        "Model_Group",
        "Final_Overall",
        "Player",
    ],
    ascending=[
        True,
        True,
        False,
        True,
    ],
    na_position="last"
).drop(
    columns="_Position_Sort"
).reset_index(drop=True)


# ==================================================
# FINAL COLUMN ORDER
# ==================================================

if KEEP_ALL_POSITION_METRICS:
    master_columns = ordered_master_columns(
        master
    )
else:
    master_columns = [
        column
        for column in COMPACT_COLUMNS
        if column in master.columns
    ]


master = master[
    master_columns
].copy()


# ==================================================
# ROUND COMMON MASTER FIELDS
# ==================================================

rounding = {
    "Age": 1,
    "Final_Overall": 2,
    "Age_Adjusted_Overall": 2,
    "Legacy_Rating": 2,
    "Legacy_Rating_Used": 2,
    "Advanced_Rating": 2,
    "Advanced_Weighted_Pctl": 2,
    "Pre_Bonus_Overall": 2,
    "Historical_Record_Bonus": 2,
    "Advanced_Data_Coverage": 3,
    "Actual_Advanced_Weight": 3,
    "Actual_Legacy_Weight": 3,
    "Madden": 1,
    "PFF": 1,
    "PFR": 1,
}

master = master.round(
    {
        column: decimals
        for column, decimals in rounding.items()
        if column in master.columns
    }
)


# ==================================================
# NORMALIZE IDENTIFIER TYPES
#
# Position Parquets often store IDs as integers while the QB Excel
# fallback can load those same IDs as strings. Pandas allows that mixed
# object column, but PyArrow does not. Normalize all ID-like columns to
# nullable strings before saving the combined master.
# ==================================================

master = normalize_identifier_columns(
    master
)


# ==================================================
# SAVE MASTER PARQUET
# ==================================================

master.to_parquet(
    MASTER_PARQUET,
    index=False
)


# ==================================================
# SAVE ONE-SHEET EXCEL MASTER
# ==================================================

try:
    with pd.ExcelWriter(
        MASTER_OUTPUT,
        engine="openpyxl"
    ) as writer:

        master.to_excel(
            writer,
            sheet_name="All_Players",
            index=False
        )

    format_master_workbook(
        MASTER_OUTPUT
    )

except PermissionError:
    raise PermissionError(
        f"Close {MASTER_OUTPUT} in Excel and run the script again."
    )


# ==================================================
# SUMMARY
# ==================================================

print(
    f"\nCreated {MASTER_OUTPUT}"
)

print(
    f"Created {MASTER_PARQUET}"
)

print(
    f"Total players: {len(master)}"
)

print(
    "\nPlayers by model:"
)

print(
    master["Model_Group"]
    .value_counts()
    .sort_index()
    .to_string()
)

print(
    "\nMaster build complete."
)
